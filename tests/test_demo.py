# -*- coding: utf-8 -*-
"""冒烟测试：把 examples/demo 跑一遍，校验「全文黑色 / 三档固定行距 / 报告表结构 / 可重复生成」。

用法：
    python tests/test_demo.py      # 直接跑（CI 用这个）
    pytest tests/test_demo.py      # 也可以用 pytest

如果当前环境的临时目录不可写（比如受限沙箱），用环境变量指定一个可写目录：
    SQL_REPORT_WORKDIR=<可写目录> python tests/test_demo.py
"""
import hashlib
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SKILL = ROOT / "skills" / "sql_report_make"
SCRIPT = SKILL / "scripts" / "make_sql_report.py"
DEMO = SKILL / "examples" / "demo"

# 全篇只允许这三种行距（固定值 exact）：正文/标题 18 磅、报告表 14 磅、代码 12 磅
LINE_OK = {("360", "exact"), ("280", "exact"), ("240", "exact")}
SECTIONS = ["1．实验目的：", "实验内容和要求：", "实验重点和难点：", "实验程序：", "实验总结："]


def run_generator(cfg, cwd):
    env = dict(os.environ, PYTHONIOENCODING="utf-8")
    proc = subprocess.run(
        [sys.executable, str(SCRIPT), str(cfg)],
        cwd=str(cwd), capture_output=True, text=True,
        encoding="utf-8", errors="replace", env=env,
    )
    assert proc.returncode == 0, f"生成失败：\n{proc.stdout}\n{proc.stderr}"
    return proc.stdout


def walk(container):
    """递归遍历段落（含嵌套表格），python-docx 的 .paragraphs 不含表格内部。"""
    for p in container.paragraphs:
        yield p
    for tb in container.tables:
        for row in tb.rows:
            seen = set()
            for cell in row.cells:
                if id(cell._tc) in seen:      # 合并单元格会重复返回
                    continue
                seen.add(id(cell._tc))
                yield from walk(cell)


def read_member(path, member):
    with zipfile.ZipFile(path) as z:
        return z.read(member)


def sha256_of_member(path, member):
    return hashlib.sha256(read_member(path, member)).hexdigest()


def check(work):
    """在 work 目录（demo 的副本）里生成并校验。"""
    cfg = work / "config-example.json"
    docx_path = work / "out" / "实验报告.docx"

    # 1) 能跑通，且 fence_to_sql 全部命中
    out = run_generator(cfg, work)
    assert "命中 2/2" in out, out
    assert docx_path.exists(), out
    assert "全文黑色" in out, out

    xml = read_member(docx_path, "word/document.xml").decode("utf-8")

    # 2) 全角空格已换成双空格
    assert "\u3000" not in xml

    # 3) 报告表：4 行 × 6 列，标题行与内容行各合并一次
    assert xml.count('w:gridSpan w:val="6"') == 2

    # 4) 默认不生成目录（无 PAGEREF 域、无书签）
    assert "PAGEREF" not in xml and "bookmarkStart" not in xml

    # 5) 行距只有三档，且都是固定值
    rules = set(re.findall(r'w:line="(\d+)" w:lineRule="(\w+)"', xml))
    assert rules <= LINE_OK, f"出现非约定的行距：{rules - LINE_OK}"
    assert {r[0] for r in rules} == {"360", "280", "240"}, rules

    # 6) 全文黑色：document.xml 里的颜色与实际 run 都必须是 000000
    assert set(re.findall(r'<w:color w:val="([^"]+)"', xml)) <= {"000000"}
    doc = Document(str(docx_path))
    for p in walk(doc):
        for r in p.runs:
            if r.font.color is not None and r.font.color.rgb is not None:
                assert str(r.font.color.rgb) == "000000", (p.text, r.text)

    # 7) 主题配色也被改成黑色
    theme = read_member(docx_path, "word/theme/theme1.xml").decode("utf-8")
    assert set(re.findall(r'<a:srgbClr val="([0-9A-Fa-f]{6})"', theme)) <= {"000000"}

    # 8) 报告表结构：4 行 × 6 列，栏目齐全，内含嵌套表格
    form = doc.tables[0]
    assert len(form.rows) == 4, len(form.rows)
    assert len(form.columns) == 6, len(form.columns)
    text = "\n".join(p.text for p in walk(doc))
    for name in SECTIONS:
        assert name in text, f"缺少栏目：{name}"
    assert sum(len(c.tables) for row in form.rows for c in row.cells) >= 1

    # 9) fence_to_sql 真的把 SQL 源文件换进来了（这行注释来自 01_create_database.sql）
    assert "创建实验用数据库" in text

    # 10) 代码块里没有装饰性横幅批注
    assert "/* ====" not in text

    # 11) 可重复生成：同一输入两次，document.xml 的 sha256 必须相同
    first = sha256_of_member(docx_path, "word/document.xml")
    run_generator(cfg, work)
    second = sha256_of_member(docx_path, "word/document.xml")
    assert first == second, "两次生成结果不同，检查是否用了 hash() 之类的随机量"


def prepare(work):
    if work.exists():
        shutil.rmtree(work)
    work.parent.mkdir(parents=True, exist_ok=True)
    shutil.copytree(DEMO, work, ignore=shutil.ignore_patterns("out", "__pycache__"))


def test_demo():
    work_root = os.environ.get("SQL_REPORT_WORKDIR")
    if work_root:
        work = Path(work_root) / "demo"
        prepare(work)
        check(work)
    else:
        with tempfile.TemporaryDirectory() as td:
            work = Path(td) / "demo"
            prepare(work)
            check(work)
    print("OK：demo 生成成功，黑色 / 行距 / 报告表 / 可重复 全部通过")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    test_demo()
